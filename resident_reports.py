"""
resident_reports.py - Resident & household registration reports for BoñgaTrack.

Companion to reports.py (which covers inventory & relief claims). This one
answers "who is registered, where are they, and where are they right now".

Setup
-----
1. pip install openpyxl reportlab
2. Put this file next to app.py
3. In app.py, right after the existing reports registration, add:

       from resident_reports import register_resident_report_routes
       register_resident_report_routes(app, get_db_connection)

Routes added (admin session required)
-------------------------------------
    GET /api/reports/residents.xlsx?purok=all|<purok>&status=all|evacuated|home
    GET /api/reports/residents.pdf?purok=...&status=...
    GET /api/reports/residents.csv?purok=...&status=...
    GET /api/reports/residents/summary      -> JSON, for on-screen preview cards
"""

import csv
import io
from datetime import datetime, timedelta, timezone
from xml.sax.saxutils import escape

from flask import jsonify, make_response, request, send_file, session

PH_TZ = timezone(timedelta(hours=8))
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_HEX = "1F3A5F"
ZEBRA_HEX = "F3F6FA"
MUTED_HEX = "5B6675"

PRIORITY_KEYWORDS = ["senior", "pwd", "disab", "pregnan"]

# Age brackets used for the CDRRMO / DSWD style breakdown.
AGE_BANDS = [
    ("Infant (0-1)", 0, 1),
    ("Child (2-12)", 2, 12),
    ("Teen (13-17)", 13, 17),
    ("Adult (18-59)", 18, 59),
    ("Senior (60+)", 60, 200),
]


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def _is_admin():
    rid = (session.get("resident_id") or "").lower()
    role = (session.get("role") or "").lower()
    return rid == "admin" or role == "admin"


def _age_band(age):
    try:
        age = int(age or 0)
    except (TypeError, ValueError):
        age = 0
    for label, lo, hi in AGE_BANDS:
        if lo <= age <= hi:
            return label
    return "Unknown"


def _is_priority(vuln):
    v = (vuln or "").strip().lower()
    if not v or v == "none":
        return False
    return any(k in v for k in PRIORITY_KEYWORDS)


def _scope_label(data):
    bits = []
    bits.append("All puroks" if data["purok_filter"] is None else data["purok_filter"])
    bits.append({"evacuated": "Evacuated households only",
                 "home": "At-home households only"}.get(data["status_filter"], "All statuses"))
    return "  |  ".join(bits)


def _filename(data, ext):
    scope = "AllPuroks" if data["purok_filter"] is None else \
        data["purok_filter"].replace(" ", "")
    return f"BongaTrack_Resident_Report_{scope}_{data['generated'][:10]}.{ext}"


# ----------------------------------------------------------------------
# Data collection
# ----------------------------------------------------------------------
def _collect(conn, purok_param, status_param):
    purok_filter = (purok_param or "all").strip()
    purok_filter = None if purok_filter.lower() in ("", "all") else purok_filter

    status_filter = (status_param or "all").strip().lower()
    if status_filter not in ("evacuated", "home"):
        status_filter = "all"

    sql = "SELECT * FROM residents WHERE id != 'admin'"
    params = []
    if purok_filter:
        sql += " AND purok = ?"
        params.append(purok_filter)
    sql += " ORDER BY purok, head_name"
    resident_rows = conn.execute(sql, params).fetchall()

    members_by_household = {}
    for m in conn.execute(
        "SELECT * FROM family_members ORDER BY resident_id, id"
    ).fetchall():
        members_by_household.setdefault(m["resident_id"], []).append(m)

    # Registration timestamps come from the transaction log.
    registered_at = {}
    for lg in conn.execute(
        "SELECT actor, time FROM transaction_log WHERE action='REGISTRATION' ORDER BY id"
    ).fetchall():
        registered_at.setdefault(lg["actor"], lg["time"])

    claims_count = {}
    last_claim = {}
    for c in conn.execute(
        "SELECT resident_id, batch_number, claim_time FROM relief_claims ORDER BY id"
    ).fetchall():
        claims_count[c["resident_id"]] = claims_count.get(c["resident_id"], 0) + 1
        last_claim[c["resident_id"]] = c["claim_time"]

    households = []
    members = []
    purok_stats = {}
    vuln_stats = {}
    age_stats = {label: 0 for label, _, _ in AGE_BANDS}

    totals = {
        "households": 0, "members": 0, "evacuated": 0, "at_home": 0,
        "vulnerable": 0, "priority_households": 0, "no_mobile": 0,
        "no_roster": 0, "gps_known": 0,
    }

    for r in resident_rows:
        roster = members_by_household.get(r["id"], [])
        evac = [m for m in roster if (m["status"] or "") == "Evacuated"]
        home = [m for m in roster if (m["status"] or "") != "Evacuated"]

        if status_filter == "evacuated" and not evac:
            continue
        if status_filter == "home" and evac:
            continue

        priority_flags = [f"{m['full_name']} ({m['vulnerability']})"
                          for m in roster if _is_priority(m["vulnerability"])]
        vulnerable = [m for m in roster
                      if (m["vulnerability"] or "None").strip().lower() not in ("", "none")]
        centers = sorted({(m["evac_center"] or "").strip() for m in evac if m["evac_center"]})

        if not roster:
            household_status = "No roster encoded"
        elif len(evac) == len(roster):
            household_status = "Fully evacuated"
        elif evac:
            household_status = f"Partially evacuated ({len(evac)}/{len(roster)})"
        else:
            household_status = "At home"

        mobile = (r["mobile"] or "").strip()
        gps = ""
        if r["current_lat"] is not None and r["current_lon"] is not None:
            gps = f"{round(r['current_lat'], 5)}, {round(r['current_lon'], 5)}"
            totals["gps_known"] += 1

        households.append({
            "id": r["id"],
            "head_name": r["head_name"],
            "purok": r["purok"] or "-",
            "mobile": mobile or "Not provided",
            "family_size": len(roster),
            "evacuated": len(evac),
            "at_home": len(home),
            "vulnerable": len(vulnerable),
            "status": household_status,
            "centers": ", ".join(centers) if centers else "-",
            "priority": "Yes - " + "; ".join(priority_flags) if priority_flags else "No",
            "registered": registered_at.get(r["id"], "-"),
            "claims": claims_count.get(r["id"], 0),
            "last_claim": last_claim.get(r["id"], "-"),
            "gps": gps or "No GPS ping yet",
            "last_active": r["last_active"] or "-",
        })

        totals["households"] += 1
        totals["members"] += len(roster)
        totals["evacuated"] += len(evac)
        totals["at_home"] += len(home)
        totals["vulnerable"] += len(vulnerable)
        if priority_flags:
            totals["priority_households"] += 1
        if not mobile or mobile.upper() == "N/A":
            totals["no_mobile"] += 1
        if not roster:
            totals["no_roster"] += 1

        p = purok_stats.setdefault(r["purok"] or "-", {
            "households": 0, "members": 0, "evacuated": 0, "at_home": 0,
            "vulnerable": 0, "priority": 0})
        p["households"] += 1
        p["members"] += len(roster)
        p["evacuated"] += len(evac)
        p["at_home"] += len(home)
        p["vulnerable"] += len(vulnerable)
        if priority_flags:
            p["priority"] += 1

        for m in roster:
            band = _age_band(m["age"])
            age_stats[band] = age_stats.get(band, 0) + 1
            v = (m["vulnerability"] or "None").strip() or "None"
            vuln_stats[v] = vuln_stats.get(v, 0) + 1
            members.append({
                "resident_id": r["id"],
                "head_name": r["head_name"],
                "purok": r["purok"] or "-",
                "full_name": m["full_name"],
                "age": m["age"] if m["age"] is not None else "-",
                "age_band": band,
                "vulnerability": v,
                "status": m["status"] or "At Home",
                "center": m["evac_center"] or "-",
            })

    center_stats = {}
    for m in members:
        if m["status"] == "Evacuated" and m["center"] != "-":
            center_stats[m["center"]] = center_stats.get(m["center"], 0) + 1

    return {
        "generated": datetime.now(PH_TZ).strftime("%Y-%m-%d %H:%M"),
        "purok_filter": purok_filter,
        "status_filter": status_filter,
        "households": households,
        "members": members,
        "totals": totals,
        "purok_stats": sorted(purok_stats.items()),
        "vuln_stats": sorted(vuln_stats.items(), key=lambda kv: -kv[1]),
        "age_stats": [(label, age_stats.get(label, 0)) for label, _, _ in AGE_BANDS],
        "center_stats": sorted(center_stats.items()),
    }


# ----------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------
def _build_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=14, bold=True)
    muted = Font(name="Arial", size=9, color=MUTED_HEX)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=HEADER_HEX)
    zebra = PatternFill("solid", fgColor=ZEBRA_HEX)
    line = Side(style="thin", color="D5DBE3")
    box = Border(left=line, right=line, top=line, bottom=line)

    wb = Workbook()
    scope = _scope_label(data)
    subtitle = f"Generated {data['generated']} (PH time)  |  Scope: {scope}"

    def prep(ws, heading, sub):
        ws["A1"] = heading
        ws["A1"].font = title_font
        ws["A2"] = sub
        ws["A2"].font = muted
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    def table(ws, headers, rows, widths, start=4, wrap_cols=()):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=start, column=i, value=h)
            c.font, c.fill, c.border = head_font, head_fill, box
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r_i, row in enumerate(rows, start + 1):
            for c_i, val in enumerate(row, 1):
                c = ws.cell(row=r_i, column=c_i, value=val)
                c.font, c.border = base, box
                c.alignment = Alignment(vertical="top", wrap_text=(c_i in wrap_cols))
                if (r_i - start) % 2 == 0:
                    c.fill = zebra
        if not rows:
            ws.cell(row=start + 1, column=1, value="No records.").font = muted
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        if rows:
            ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"

    def mini(ws, row, heading, headers, rows):
        ws.cell(row=row, column=1, value=heading).font = bold
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row + 1, column=i, value=h)
            c.font, c.fill, c.border = head_font, head_fill, box
        for j, values in enumerate(rows, row + 2):
            for i, val in enumerate(values, 1):
                c = ws.cell(row=j, column=i, value=val)
                c.font, c.border = base, box
                c.alignment = Alignment(horizontal="left")
        return row + 2 + len(rows) + 1

    t = data["totals"]

    # ---- Summary ------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    prep(ws, "BoñgaTrack - Resident & Household Registration Report", subtitle)
    summary = [
        ("Registered households", t["households"]),
        ("Registered individuals", t["members"]),
        ("Currently evacuated", t["evacuated"]),
        ("Currently at home", t["at_home"]),
        ("Individuals with a vulnerability", t["vulnerable"]),
        ("Priority households (senior / PWD / pregnant)", t["priority_households"]),
        ("Households with no mobile number", t["no_mobile"]),
        ("Households with an empty roster", t["no_roster"]),
        ("Households with a GPS ping on record", t["gps_known"]),
    ]
    for i, (k, v) in enumerate(summary, 4):
        a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=v)
        a.font, b.font = bold, base
        a.border = b.border = box
        b.alignment = Alignment(horizontal="left")

    r = 4 + len(summary) + 1
    r = mini(ws, r, "Age profile",
             ["Age group", "Individuals"], [[k, v] for k, v in data["age_stats"]])
    r = mini(ws, r, "Vulnerability profile",
             ["Vulnerability", "Individuals"], [[k, v] for k, v in data["vuln_stats"]])
    mini(ws, r, "Head count per evacuation centre",
         ["Evacuation centre", "Individuals"],
         [[k, v] for k, v in data["center_stats"]] or [["-", 0]])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 22

    # ---- Purok breakdown ---------------------------------------------
    ws = wb.create_sheet("Purok Breakdown")
    prep(ws, "Registration by Purok", subtitle)
    table(ws, ["Purok", "Households", "Individuals", "Evacuated", "At home",
               "Vulnerable", "Priority households"],
          [[p, s["households"], s["members"], s["evacuated"], s["at_home"],
            s["vulnerable"], s["priority"]] for p, s in data["purok_stats"]],
          [18, 14, 14, 14, 12, 14, 20])

    # ---- Household directory -----------------------------------------
    ws = wb.create_sheet("Household Directory")
    prep(ws, "Household Directory", subtitle)
    table(ws, ["Household ID", "Household Head", "Purok", "Mobile", "Family size",
               "Evacuated", "At home", "Vulnerable", "Household status",
               "Evacuation centre", "Priority household", "Registered on",
               "Relief claims", "Last claim", "Last known GPS", "Last active"],
          [[h["id"], h["head_name"], h["purok"], h["mobile"], h["family_size"],
            h["evacuated"], h["at_home"], h["vulnerable"], h["status"], h["centers"],
            h["priority"], h["registered"], h["claims"], h["last_claim"],
            h["gps"], h["last_active"]] for h in data["households"]],
          [16, 28, 12, 15, 12, 12, 10, 12, 24, 30, 34, 20, 13, 20, 24, 20],
          wrap_cols=(10, 11))

    # ---- Individual roster -------------------------------------------
    ws = wb.create_sheet("Individual Roster")
    prep(ws, "Individual Roster (every registered person)", subtitle)
    table(ws, ["Household ID", "Household Head", "Purok", "Full name", "Age",
               "Age group", "Vulnerability", "Status", "Evacuation centre"],
          [[m["resident_id"], m["head_name"], m["purok"], m["full_name"], m["age"],
            m["age_band"], m["vulnerability"], m["status"], m["center"]]
           for m in data["members"]],
          [16, 28, 12, 28, 8, 16, 22, 14, 30])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def _build_pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Table,
                                    TableStyle)

    header_c = colors.HexColor("#" + HEADER_HEX)
    zebra_c = colors.HexColor("#" + ZEBRA_HEX)
    muted_c = colors.HexColor("#" + MUTED_HEX)
    grid_c = colors.HexColor("#D5DBE3")

    ss = getSampleStyleSheet()
    st = {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold",
                                fontSize=18, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=9,
                              textColor=muted_c, spaceAfter=10),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold",
                             fontSize=12, spaceBefore=14, spaceAfter=6, keepWithNext=1),
        "th": ParagraphStyle("th", parent=ss["Normal"], fontName="Helvetica-Bold",
                             fontSize=8.5, textColor=colors.white, leading=10),
        "td": ParagraphStyle("td", parent=ss["Normal"], fontSize=8.5, leading=10.5),
        "tdr": ParagraphStyle("tdr", parent=ss["Normal"], fontSize=8.5,
                              leading=10.5, alignment=TA_RIGHT),
        "note": ParagraphStyle("n", parent=ss["Normal"], fontSize=9, textColor=muted_c),
    }

    def P(text, style):
        return Paragraph(escape(str(text)), st[style])

    def make_table(headers, rows, widths, right_cols=()):
        head = [P(h, "th") for h in headers]
        body = [[P(v, "tdr" if i in right_cols else "td") for i, v in enumerate(r)]
                for r in rows]
        t = Table([head] + body, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), header_c),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, zebra_c]),
            ("GRID", (0, 0), (-1, -1), 0.4, grid_c),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    def section(title, headers, rows, widths, right_cols=()):
        out = [Paragraph(escape(title), st["h2"])]
        out.append(make_table(headers, rows, widths, right_cols) if rows
                   else Paragraph("No records.", st["note"]))
        return out

    page = landscape(A4)
    margin = 15 * mm
    width = page[0] - 2 * margin
    scope = _scope_label(data)
    t = data["totals"]

    story = [
        Paragraph("BoñgaTrack - Resident &amp; Household Registration Report", st["title"]),
        Paragraph(escape(f"Generated {data['generated']} (PH time)  |  Scope: {scope}"), st["sub"]),
    ]

    left = make_table(
        ["Summary", ""],
        [["Registered households", t["households"]],
         ["Registered individuals", t["members"]],
         ["Currently evacuated", t["evacuated"]],
         ["Currently at home", t["at_home"]],
         ["With a vulnerability", t["vulnerable"]],
         ["Priority households", t["priority_households"]],
         ["No mobile number on file", t["no_mobile"]],
         ["Empty roster", t["no_roster"]]],
        [width * 0.30, width * 0.14], right_cols=(1,))
    right = make_table(
        ["Age group", "Individuals"],
        [[k, v] for k, v in data["age_stats"]],
        [width * 0.22, width * 0.14], right_cols=(1,))
    wrapper = Table([[left, right]], colWidths=[width * 0.5, width * 0.5])
    wrapper.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(wrapper)

    story += section(
        "Registration by Purok",
        ["Purok", "Households", "Individuals", "Evacuated", "At home",
         "Vulnerable", "Priority"],
        [[p, s["households"], s["members"], s["evacuated"], s["at_home"],
          s["vulnerable"], s["priority"]] for p, s in data["purok_stats"]],
        [width * 0.16] + [width * 0.14] * 6, right_cols=(1, 2, 3, 4, 5, 6))

    story += section(
        "Head count per evacuation centre",
        ["Evacuation centre", "Individuals"],
        [[k, v] for k, v in data["center_stats"]],
        [width * 0.6, width * 0.2], right_cols=(1,))

    story += section(
        "Household Directory",
        ["Household ID", "Household Head", "Purok", "Mobile", "Size", "Evac.",
         "Vuln.", "Status", "Evacuation centre", "Priority", "Registered on"],
        [[h["id"], h["head_name"], h["purok"], h["mobile"], h["family_size"],
          h["evacuated"], h["vulnerable"], h["status"], h["centers"],
          h["priority"], h["registered"]] for h in data["households"]],
        [width * 0.08, width * 0.13, width * 0.06, width * 0.08, width * 0.04,
         width * 0.04, width * 0.04, width * 0.11, width * 0.14, width * 0.16,
         width * 0.10], right_cols=(4, 5, 6))

    story += section(
        "Individual Roster",
        ["Household ID", "Full name", "Age", "Age group", "Vulnerability",
         "Status", "Evacuation centre"],
        [[m["resident_id"], m["full_name"], m["age"], m["age_band"],
          m["vulnerability"], m["status"], m["center"]] for m in data["members"]],
        [width * 0.11, width * 0.21, width * 0.06, width * 0.13, width * 0.16,
         width * 0.11, width * 0.22], right_cols=(2,))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted_c)
        canvas.drawString(margin, 9 * mm,
                          f"BoñgaTrack - generated {data['generated']}")
        canvas.drawRightString(page[0] - margin, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page, leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=16 * mm,
        title="BoñgaTrack Resident Registration Report", author="BoñgaTrack")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


# ----------------------------------------------------------------------
# CSV (lightweight, opens anywhere, no extra libraries)
# ----------------------------------------------------------------------
def _build_csv(data):
    out = io.StringIO()
    w = csv.writer(out)
    t = data["totals"]

    w.writerow(["BoñgaTrack - Resident & Household Registration Report"])
    w.writerow([f"Generated {data['generated']} (PH time)", f"Scope: {_scope_label(data)}"])
    w.writerow([])
    w.writerow(["SUMMARY"])
    for k, v in [("Registered households", t["households"]),
                 ("Registered individuals", t["members"]),
                 ("Currently evacuated", t["evacuated"]),
                 ("Currently at home", t["at_home"]),
                 ("With a vulnerability", t["vulnerable"]),
                 ("Priority households", t["priority_households"]),
                 ("No mobile number on file", t["no_mobile"]),
                 ("Empty roster", t["no_roster"])]:
        w.writerow([k, v])

    w.writerow([])
    w.writerow(["BY PUROK"])
    w.writerow(["Purok", "Households", "Individuals", "Evacuated", "At home",
                "Vulnerable", "Priority households"])
    for p, s in data["purok_stats"]:
        w.writerow([p, s["households"], s["members"], s["evacuated"],
                    s["at_home"], s["vulnerable"], s["priority"]])

    w.writerow([])
    w.writerow(["HOUSEHOLD DIRECTORY"])
    w.writerow(["Household ID", "Household Head", "Purok", "Mobile", "Family size",
                "Evacuated", "At home", "Vulnerable", "Household status",
                "Evacuation centre", "Priority household", "Registered on",
                "Relief claims", "Last claim", "Last known GPS", "Last active"])
    for h in data["households"]:
        w.writerow([h["id"], h["head_name"], h["purok"], h["mobile"], h["family_size"],
                    h["evacuated"], h["at_home"], h["vulnerable"], h["status"],
                    h["centers"], h["priority"], h["registered"], h["claims"],
                    h["last_claim"], h["gps"], h["last_active"]])

    w.writerow([])
    w.writerow(["INDIVIDUAL ROSTER"])
    w.writerow(["Household ID", "Household Head", "Purok", "Full name", "Age",
                "Age group", "Vulnerability", "Status", "Evacuation centre"])
    for m in data["members"]:
        w.writerow([m["resident_id"], m["head_name"], m["purok"], m["full_name"],
                    m["age"], m["age_band"], m["vulnerability"], m["status"],
                    m["center"]])

    return out.getvalue()


# ----------------------------------------------------------------------
# Routes
# ----------------------------------------------------------------------
def register_resident_report_routes(app, get_db_connection):

    def _load():
        conn = get_db_connection()
        try:
            return _collect(conn, request.args.get("purok"), request.args.get("status"))
        finally:
            conn.close()

    def _guard():
        if not session.get("resident_id"):
            return jsonify({"status": "error", "message": "Please log in first."}), 401
        if not _is_admin():
            return jsonify({"status": "error", "message": "Admin access only."}), 403
        return None

    def build(builder, ext, mimetype):
        blocked = _guard()
        if blocked:
            return blocked
        data = _load()
        try:
            buf = builder(data)
        except ImportError as e:
            return jsonify({
                "status": "error",
                "message": f"Missing library on the server: {e.name}. "
                           f"Run: pip install openpyxl reportlab"
            }), 500
        return send_file(buf, as_attachment=True,
                         download_name=_filename(data, ext), mimetype=mimetype)

    @app.route("/api/reports/residents.xlsx")
    def report_residents_xlsx():
        return build(_build_xlsx, "xlsx", XLSX_MIME)

    @app.route("/api/reports/residents.pdf")
    def report_residents_pdf():
        return build(_build_pdf, "pdf", "application/pdf")

    @app.route("/api/reports/residents.csv")
    def report_residents_csv():
        blocked = _guard()
        if blocked:
            return blocked
        data = _load()
        resp = make_response(_build_csv(data))
        resp.headers["Content-Disposition"] = \
            f"attachment; filename={_filename(data, 'csv')}"
        resp.headers["Content-Type"] = "text/csv; charset=utf-8"
        return resp

    @app.route("/api/reports/residents/summary")
    def report_residents_summary():
        blocked = _guard()
        if blocked:
            return blocked
        data = _load()
        return jsonify({
            "status": "success",
            "generated": data["generated"],
            "scope": _scope_label(data),
            "totals": data["totals"],
            "purok_stats": [{"purok": p, **s} for p, s in data["purok_stats"]],
            "age_stats": [{"band": k, "count": v} for k, v in data["age_stats"]],
            "vuln_stats": [{"vulnerability": k, "count": v} for k, v in data["vuln_stats"]],
            "center_stats": [{"center": k, "count": v} for k, v in data["center_stats"]],
        })